"""Alzea: consulta de datos y análisis.

Acciones: Yahoo Finance (datos de mercado). Deuda, tipo de cambio y tasas: Banco de México (SIE).
Entorno de EE. UU.: Tesoro de EE. UU., Reserva Federal de Nueva York, BLS y Cboe (ver fuentes.py).
En producción lo usan las funciones de Vercel (api/*.py); en la Mac puede correr como servidor
local (solo 127.0.0.1) con:  python server.py
"""
import io
import json
import re
import threading
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from zoneinfo import ZoneInfo

import analisis
import diario
import reportes
import truststore
from openpyxl import load_workbook

truststore.inject_into_ssl()  # usa los certificados de macOS/Windows (redes con proxy)

PORT = 8765
TZ = ZoneInfo("America/Mexico_City")
STATIC = Path(__file__).parent / "static"
UA = {"User-Agent": "Mozilla/5.0 (Macintosh) Alzea/1.0"}

# Primer día consultable: todos los periodos arrancan el 1 de marzo de 2026 o después.
INICIO = date(2026, 3, 1)

ACCIONES = {
    "WALMEX.MX": "Walmart de México",
    "AMXB.MX": "América Móvil",
    "KOFUBL.MX": "Coca-Cola FEMSA",
    "KIMBERA.MX": "Kimberly-Clark de México",
    "OMAB.MX": "Grupo Aeroportuario Centro Norte",
    "AAPL": "Apple",
    "MSFT": "Microsoft",
    "NVDA": "NVIDIA",
    "TSLA": "Tesla",
    "KO": "The Coca-Cola Company",
}
META = {
    "WALMEX.MX": {"mercado": "BMV", "moneda": "MXN", "pais": "México"},
    "AMXB.MX": {"mercado": "BMV", "moneda": "MXN", "pais": "México"},
    "KOFUBL.MX": {"mercado": "BMV", "moneda": "MXN", "pais": "México"},
    "KIMBERA.MX": {"mercado": "BMV", "moneda": "MXN", "pais": "México"},
    "OMAB.MX": {"mercado": "BMV", "moneda": "MXN", "pais": "México"},
    "AAPL": {"mercado": "Nasdaq", "moneda": "USD", "pais": "EE. UU."},
    "MSFT": {"mercado": "Nasdaq", "moneda": "USD", "pais": "EE. UU."},
    "NVDA": {"mercado": "Nasdaq", "moneda": "USD", "pais": "EE. UU."},
    "TSLA": {"mercado": "Nasdaq", "moneda": "USD", "pais": "EE. UU."},
    "KO": {"mercado": "NYSE", "moneda": "USD", "pais": "EE. UU."},
}

# cuadro Banxico, sección, renglón principal, emisor, unidad, tipo, renglones complementarios
DEUDA = {
    "CETES 28 días": dict(codigo="CETES 28", cuadro="CF107", seccion="Cetes a 28 días", renglon="Tasa de rendimiento",
                          emisor="Gobierno Federal", unidad="% anual", tipo="tasa",
                          extras=[("Monto asignado (millones de pesos)", "Monto asignado")]),
    "BONOS M 10 años": dict(codigo="BONO M 10A", cuadro="CF107", seccion="Bonos a tasa fija a 10 años", renglon="Tasa de rendimiento",
                            emisor="Gobierno Federal", unidad="% anual", tipo="tasa",
                            extras=[("Monto asignado (millones de pesos)", "Monto asignado")]),
    "UDIBONOS 10 años": dict(codigo="UDIBONO 10A", cuadro="CF107", seccion="Udibonos a 10 años", renglon="Tasa de rendimiento real",
                             emisor="Gobierno Federal", unidad="% real anual", tipo="tasa",
                             extras=[("Monto asignado (millones de UDIS)", "Monto asignado")]),
    "PAPEL COMERCIAL": dict(codigo="PAPEL COM.", cuadro="CF133", seccion="Papel Comercial", renglon="Monto colocado",
                            emisor="Empresas emisoras (BMV)", unidad="miles de pesos colocados por semana", tipo="monto",
                            extras=[("Saldo vigente (miles de pesos)", "Saldo vigente")]),
    "CERTIFICADOS BURSÁTILES": dict(codigo="CEBURES CP", cuadro="CF133", seccion="Certificados Bursátiles a Corto Plazo",
                                    renglon="Tasa promedio ponderada", emisor="Empresas emisoras (BMV)",
                                    unidad="% anual (corto plazo)", tipo="tasa",
                                    extras=[("Monto colocado (miles de pesos)", "Monto colocado"),
                                            ("Saldo vigente (miles de pesos)", "Saldo vigente"),
                                            ("Plazo promedio de colocación (días)", "Plazo promedio ponderado de colocación")]),
    "CETES 91 días (referencia)": dict(codigo="CETES 91", cuadro="CF107", seccion="Cetes a 91 días", renglon="Tasa de rendimiento",
                                       emisor="Gobierno Federal", unidad="% anual", tipo="tasa", extras=[], oculto=True),
    "BONDES F 2 años": dict(codigo="BONDE F 2A", cuadro="CF107", seccion="Bondes F a 2 años", renglon="Precio Promedio",
                            emisor="Gobierno Federal", unidad="$ por título de $100", tipo="precio",
                            extras=[("Monto asignado (millones de pesos)", "Monto asignado")]),
    "BPAG28 3 años": dict(codigo="BPAG28 3A", cuadro="CF115", seccion="BPAG28", renglon="Sobretasa",
                          emisor="IPAB", unidad="puntos % sobre referencia", tipo="tasa",
                          extras=[("Monto asignado (millones de pesos)", "Monto asignado")]),
}

_cache, _lock = {}, threading.Lock()


def cached(key, ttl, fn):
    with _lock:
        hit = _cache.get(key)
        if hit and time.time() - hit[0] < ttl:
            return hit[1]
    val = fn()
    with _lock:
        _cache[key] = (time.time(), val)
    return val


def http_get(url, timeout=60, reintentos=3, cabeceras=None):
    """GET con reintentos (los portales oficiales limitan ráfagas con HTTP 429)."""
    for k in range(reintentos):
        try:
            req = urllib.request.Request(url, headers={**UA, **(cabeceras or {})})
            with urllib.request.urlopen(req, timeout=timeout) as r:
                return r.read()
        except urllib.error.HTTPError as e:
            if e.code in (429, 500, 502, 503, 504) and k < reintentos - 1:
                time.sleep(1.5 * (k + 1))
                continue
            raise
        except (TimeoutError, urllib.error.URLError):
            if k < reintentos - 1:
                time.sleep(1.0)
                continue
            raise


def sin_acentos(t):
    return "".join(c for c in unicodedata.normalize("NFKD", str(t)) if not unicodedata.combining(c)).lower()


def mercado_abierto(now=None):
    # La BMV sigue el horario de Nueva York (9:30–16:00 ET): 7:30–14:00 u 8:30–15:00 en CDMX según el horario de verano de EE. UU.
    ny = (now or datetime.now(TZ)).astimezone(ZoneInfo("America/New_York"))
    return ny.weekday() < 5 and (9, 30) <= (ny.hour, ny.minute) < (16, 0)


# ---------------- Yahoo Finance ----------------
def yahoo_chart(ticker, **params):
    q = urllib.parse.urlencode(params)
    last = None
    for host in ("query1", "query2"):
        try:
            url = f"https://{host}.finance.yahoo.com/v8/finance/chart/{urllib.parse.quote(ticker)}?{q}"
            return json.loads(http_get(url, 30))["chart"]["result"][0]
        except Exception as e:  # reintenta con el otro host
            last = e
    raise last


def historico(ticker, desde):
    p1 = int(datetime.combine(desde, datetime.min.time(), TZ).timestamp())
    r = yahoo_chart(ticker, period1=p1, period2=int(time.time()) + 3600, interval="1d")
    q, meta = r["indicators"]["quote"][0], r["meta"]
    filas = []
    for i, ts in enumerate(r.get("timestamp") or []):
        if q["close"][i] is None:
            continue
        filas.append({
            "fecha": datetime.fromtimestamp(ts, TZ).date().isoformat(),
            "apertura": q["open"][i], "maximo": q["high"][i], "minimo": q["low"][i],
            "cierre": q["close"][i], "volumen": q["volume"][i] or 0,
        })
    precio = meta.get("regularMarketPrice")
    hora = meta.get("regularMarketTime")
    if filas and precio is not None and hora:
        hoy = datetime.fromtimestamp(hora, TZ).date().isoformat()
        if filas[-1]["fecha"] == hoy:
            filas[-1]["cierre"] = precio
    for i, f in enumerate(filas):
        prev = filas[i - 1]["cierre"] if i else None
        f["var"] = f["cierre"] - prev if prev else None
        f["var_pct"] = f["cierre"] / prev - 1 if prev else None
        f["var_acum_pct"] = f["cierre"] / filas[0]["cierre"] - 1
    return {
        "ticker": ticker,
        "nombre": ACCIONES.get(ticker, meta.get("longName") or meta.get("shortName") or ticker),
        "moneda": META.get(ticker, {}).get("moneda") or meta.get("currency"),
        "mercado": META.get(ticker, {}).get("mercado"),
        "pais": META.get(ticker, {}).get("pais"),
        "precio": precio,
        "hora": datetime.fromtimestamp(hora, TZ).isoformat() if hora else None,
        "historico": filas,
    }


def intradia(ticker):
    r = yahoo_chart(ticker, range="1d", interval="5m")
    q = r["indicators"]["quote"][0]
    return [
        {"hora": datetime.fromtimestamp(ts, TZ).strftime("%H:%M"), "precio": q["close"][i]}
        for i, ts in enumerate(r.get("timestamp") or []) if q["close"][i] is not None
    ]


# ---------------- Banco de México ----------------
_MESES3 = {"ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6, "jul": 7, "ago": 8, "sep": 9, "oct": 10, "nov": 11, "dic": 12}


def _fecha_columna(v):
    """Encabezado de columna del SIE: dd/mm/aaaa (diaria o semanal) o «Mar 2026» (mensual, primer día del mes)."""
    if not isinstance(v, str):
        return None
    v = v.strip()
    if re.fullmatch(r"\d\d/\d\d/\d{4}", v):
        return datetime.strptime(v, "%d/%m/%Y").date().isoformat()
    m = re.fullmatch(r"([A-Za-zñÑ]{3})[a-z]* (\d{4})", v)
    if m and m.group(1).lower() in _MESES3:
        return date(int(m.group(2)), _MESES3[m.group(1).lower()], 1).isoformat()
    return None


def _banxico_cuadro(cuadro, desde, hasta):
    ms = lambda d: int(datetime.combine(d, datetime.min.time(), TZ).timestamp() * 1000)
    url = ("https://www.banxico.org.mx/SieInternet/consultarDirectorioInternetAction.do?"
           f"sector=22&accion=consultarCuadro&idCuadro={cuadro}&locale=es&formatoXLS.x=1"
           f"&fechaInicio={ms(desde)}&fechaFin={ms(hasta)}")
    raw = http_get(url, 120)
    if not raw.startswith(b"PK"):
        raise RuntimeError(f"Banxico no devolvió el cuadro {cuadro}")
    ws = load_workbook(io.BytesIO(raw)).active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    fila = next((r for r in rows if sum(1 for v in r[2:] if _fecha_columna(v)) >= 1), None)
    if fila is None:
        raise RuntimeError(f"El cuadro {cuadro} no trae fechas en el periodo")
    return [_fecha_columna(v) if j >= 2 else None for j, v in enumerate(fila)], rows


def banxico_cuadro(cuadro, desde, hasta):
    return cached(("bx", cuadro, desde, hasta), 3 * 3600, lambda: _banxico_cuadro(cuadro, desde, hasta))


def valores_en_seccion(fechas, rows, seccion, renglon):
    """{fecha: valor} del primer renglón que contiene `renglon` dentro de la sección `seccion` (marcada con ●)."""
    actual = None
    for r in rows:
        label = (r[1] or "") if len(r) > 1 else ""
        if "●" in label:
            actual = label.split("●")[-1].strip()
        elif actual and actual.startswith(seccion) and renglon.lower() in label.lower():
            vals = {fechas[j]: float(v) for j, v in enumerate(r) if j >= 2 and fechas[j] and isinstance(v, (int, float))}
            if vals:  # los encabezados de subsección (sin cifras) se saltan
                return vals
    return {}


def serie_deuda(nombre, cuadros):
    d = DEUDA[nombre]
    fechas, rows = cuadros[d["cuadro"]]
    principal = valores_en_seccion(fechas, rows, d["seccion"], d["renglon"])
    extras = {etq: valores_en_seccion(fechas, rows, d["seccion"], ren) for etq, ren in d["extras"]}
    datos = [{"fecha": f, "valor": v, "extra": {e: extras[e].get(f) for e in extras}} for f, v in sorted(principal.items())]
    for i, x in enumerate(datos):
        prev = datos[i - 1]["valor"] if i else None
        x["var"] = x["valor"] - prev if prev is not None else None
        x["var_pb"] = (x["valor"] - prev) * 100 if prev is not None and d["tipo"] == "tasa" else None
        x["var_pct"] = x["valor"] / prev - 1 if prev else None
    return {"nombre": nombre, "codigo": d["codigo"], "emisor": d["emisor"], "unidad": d["unidad"], "tipo": d["tipo"],
            "cuadro": d["cuadro"], "extras": [e for e, _ in d["extras"]], "datos": datos}


def deuda(desde):
    hasta = datetime.now(TZ).date()

    def cargar():
        cuadros = {}
        for c in sorted({v["cuadro"] for v in DEUDA.values()}):  # 3 cuadros, uno tras otro: Banxico limita ráfagas
            cuadros[c] = banxico_cuadro(c, desde, hasta)
        return [serie_deuda(n, cuadros) for n in DEUDA]

    # las subastas y colocaciones son semanales: basta refrescar cada 3 h
    return cached(("deuda", desde, hasta), 3 * 3600, cargar)


def acciones(tickers, desde):
    ttl = 60 if mercado_abierto() else 15 * 60

    def uno(t):
        try:
            return cached(("acc", t, desde), ttl, lambda: historico(t, desde))
        except Exception as e:
            return {"ticker": t, "nombre": ACCIONES.get(t, t), "error": str(e), "historico": []}

    with ThreadPoolExecutor(max_workers=6) as ex:
        return list(ex.map(uno, tickers))


def periodo(qs):
    """Periodo consultable: nunca antes del 1 de marzo de 2026 ni después de hoy."""
    hoy = datetime.now(TZ).date()
    d = date.fromisoformat(qs.get("desde", [INICIO.isoformat()])[0])
    h = date.fromisoformat(qs.get("hasta", [hoy.isoformat()])[0])
    d, h = max(d, INICIO), min(h, hoy)
    if d >= h:
        raise ValueError("El periodo debe empezar el 1 de marzo de 2026 o después, y «Desde» debe ser anterior a «Hasta».")
    return d, h


def _cortar(rows, hasta):
    return [r for r in rows if r["fecha"] <= hasta.isoformat()]


def _verificar_sesiones(rows, desde):
    if sum(1 for r in rows if r["fecha"] >= desde.isoformat()) < 2:
        raise ValueError("El periodo elegido no incluye al menos dos sesiones de mercado. Amplía las fechas.")


def ventana_historial():
    """Historial largo (≈14 meses) para que las medias y el backtest tengan datos."""
    return date.today() - timedelta(days=420)


def cotizacion(t):
    """Último precio y puntos del día (5 min) de una emisora."""
    r = yahoo_chart(t, range="1d", interval="5m")
    m, q = r["meta"], r["indicators"]["quote"][0]
    precio, previo = m.get("regularMarketPrice"), m.get("chartPreviousClose")
    hora = m.get("regularMarketTime")
    return {
        "ticker": t, "precio": precio, "previo": previo,
        "var": precio - previo if precio is not None and previo else None,
        "var_pct": (precio / previo - 1) * 100 if precio is not None and previo else None,
        "maximo": m.get("regularMarketDayHigh"), "minimo": m.get("regularMarketDayLow"),
        "volumen": m.get("regularMarketVolume"),
        "hora": datetime.fromtimestamp(hora, TZ).isoformat() if hora else None,
        "puntos": [round(x, 2) for x in (q.get("close") or []) if x is not None],
    }


def cotizaciones(tickers):
    def uno(t):
        try:
            return cached(("cot", t), 20, lambda: cotizacion(t))
        except Exception as e:
            return {"ticker": t, "error": str(e)}

    with ThreadPoolExecutor(max_workers=6) as ex:
        return list(ex.map(uno, tickers))


def analisis_completo(tickers, desde, hasta):
    """Decisiones, estadística y series para la pantalla, con datos hasta el último cierre del periodo."""
    accs = acciones(tickers, ventana_historial())
    items, ult = [], None
    for a in accs:
        rows = _cortar(a.get("historico", []), hasta)
        if a.get("error") or len(rows) < 30:
            items.append({"ticker": a["ticker"], "nombre": a["nombre"], "error": a.get("error") or "Historial insuficiente"})
            continue
        _verificar_sesiones(rows, desde)
        d = analisis.analizar_serie(a["nombre"], rows, desde.isoformat())
        vivo = hasta >= date.today() and rows[-1]["fecha"] == a["historico"][-1]["fecha"]
        d.update(ticker=a["ticker"], nombre=a["nombre"], moneda=a.get("moneda"), mercado=a.get("mercado"), pais=a.get("pais"),
                 precio=a.get("precio") if vivo else None, hora=a.get("hora") if vivo else None)
        items.append(d)
        ult = max(ult or d["fecha"], d["fecha"])
    ok = [x for x in items if "error" not in x]
    deu = []
    todas = deuda(date.today() - timedelta(days=300))
    cetes91 = next((x for x in todas if x["nombre"].startswith("CETES 91")), None)
    for sd in todas:
        if DEUDA[sd["nombre"]].get("oculto"):
            continue
        d = analisis.analizar_deuda(sd, hasta=hasta.isoformat(), desde=desde.isoformat())
        datos = [x for x in sd["datos"] if desde.isoformat() <= x["fecha"] <= hasta.isoformat()]
        if sd["nombre"] == "CERTIFICADOS BURSÁTILES" and cetes91 and datos:
            ref = [(x["fecha"], x["valor"]) for x in cetes91["datos"]]
            fx_ = [f for f, _ in ref]
            for x in datos:
                i = max(0, __import__("bisect").bisect_right(fx_, x["fecha"]) - 1)
                x["extra"]["Prima sobre CETES a 91 días (puntos base)"] = round((x["valor"] - ref[i][1]) * 100, 2) if ref else None
            if datos[-1]["extra"].get("Prima sobre CETES a 91 días (puntos base)") is not None:
                pr = datos[-1]["extra"]["Prima sobre CETES a 91 días (puntos base)"]
                d["por_que"].append(f"Los certificados bursátiles de corto plazo pagaron {pr:+.2f} pb sobre los CETES a 91 días en la última semana: es la prima por riesgo de crédito.")
                d["extras"] = d.get("extras", []) + ["Prima sobre CETES a 91 días (puntos base)"]
        d["datos"] = datos
        deu.append(d)
    return {"ahora": datetime.now(TZ).isoformat(), "mercado_abierto": mercado_abierto(),
            "periodo": {"desde": desde.isoformat(), "hasta": hasta.isoformat(), "inicio": INICIO.isoformat(),
                        "ultimo_cierre": ult, "vivo": hasta >= date.today()},
            "mercado": analisis.resumen_mercado(ok), "acciones": items, "deuda": deu}


def _historiales(tickers=None):
    accs = acciones(list(tickers or ACCIONES), ventana_historial())
    return {a["ticker"]: a["historico"] for a in accs if not a.get("error") and a["historico"]}


def _fx():
    """Tipo de cambio de cierre de jornada (Banxico) para convertir cierres de acciones de EE. UU. a pesos."""
    import fuentes
    return fuentes.fx_cierre(INICIO - timedelta(days=15), datetime.now(TZ).date())


def dia(fecha_iso):
    """Registro de un día: decisiones, porqué y qué pasó después (para descargar en PDF, Excel o Markdown)."""
    f = date.fromisoformat(fecha_iso)
    if f < INICIO:
        raise ValueError("Solo hay análisis desde el 1 de marzo de 2026.")
    hist = _historiales()
    reg = diario.registro_dia(fecha_iso, hist, deuda(date.today() - timedelta(days=300)), ACCIONES, META)
    if not reg:
        raise ValueError("No hubo sesión de mercado en esa fecha o no hay historial suficiente")
    reg = diario.completar([reg], hist)[0]
    try:
        import contexto
        reg["entorno"] = contexto.entorno_del_dia(fecha_iso)
    except Exception as e:  # el reporte sale aunque alguna fuente externa falle
        reg["entorno"] = {"error": str(e), "indicadores": []}
    return reg


def archivo(desde, hasta):
    """Decisiones día por día dentro del periodo y qué tan seguido acertaron (cálculo directo, sin archivos)."""
    return diario.resumen_periodo(_historiales(), desde.isoformat(), hasta.isoformat(), META)


def simulacion(desde, hasta, moneda="mxn"):
    hist = _historiales()
    fx = _fx() if moneda == "mxn" else None
    return analisis.simular(hist, desde.isoformat(), hasta.isoformat(), fx=fx, meta=META)


def comparacion(tickers, desde, hasta, moneda="mxn"):
    hist = _historiales(tickers)
    fx = _fx() if moneda == "mxn" else None
    for t in list(hist):
        _verificar_sesiones(_cortar(hist[t], hasta), desde)
    return analisis.comparar(hist, desde.isoformat(), hasta.isoformat(), fx=fx, meta=META)


def informe(desde, hasta):
    """Todo lo necesario para el informe del periodo (pantalla, PDF y Excel)."""
    import contexto
    an = analisis_completo(list(ACCIONES), desde, hasta)
    return {"periodo": an["periodo"], "analisis": an,
            "contexto": contexto.construir(desde, hasta, an["acciones"]),
            "archivo": archivo(desde, hasta), "simulacion": simulacion(desde, hasta, "mxn")}


# ---------------- Servidor HTTP ----------------
class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *a, **kw):
        super().__init__(*a, directory=str(STATIC), **kw)

    def translate_path(self, path):  # mismas rutas que en Vercel: /static/app.js
        if path.startswith("/static/"):
            path = path[len("/static"):]
        return super().translate_path(path)

    def end_headers(self):
        if not self.path.startswith("/api/"):
            self.send_header("Cache-Control", "no-cache")  # el celular siempre recibe la versión más reciente
        super().end_headers()

    def log_message(self, fmt, *args):
        if "/api/" in (args[0] if args else ""):
            print(f"[{datetime.now(TZ):%H:%M:%S}] {args[0]}")

    def send(self, body, ctype="application/json; charset=utf-8", status=200, extra=None):
        self.send_response(status)
        self.send_header("Content-Type", ctype)
        self.send_header("Cache-Control", "no-store")
        for k, v in (extra or {}).items():
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        u = urllib.parse.urlparse(self.path)
        if not u.path.startswith("/api/"):
            return super().do_GET()
        qs = urllib.parse.parse_qs(u.query)
        if "fresco" in qs:  # actualizar: vuelve a pedir precios; los datos de Banxico y del Tesoro (cambian por semana o día) conservan su caché
            with _lock:
                for k in [k for k in _cache if isinstance(k, tuple) and k[0] in ("acc", "cot", "intra", "sim", "arch", "dia", "inf", "ctx")]:
                    del _cache[k]
        try:
            tickers = [t.strip() for t in qs.get("tickers", [",".join(ACCIONES)])[0].split(",") if t.strip() in ACCIONES]
            ruta = u.path
            if ruta == "/api/estado":
                data = {"ahora": datetime.now(TZ).isoformat(), "mercado_abierto": mercado_abierto(), "inicio": INICIO.isoformat(),
                        "acciones": ACCIONES, "meta": META, "deuda": [n for n, d in DEUDA.items() if not d.get("oculto")]}
            elif ruta == "/api/intradia":
                t = qs["ticker"][0]
                if t not in ACCIONES:
                    raise ValueError("Emisora desconocida")
                data = cached(("intra", t), 60, lambda: intradia(t))
            elif ruta == "/api/cotizaciones":
                data = cotizaciones(tickers)
            elif ruta == "/api/guia":
                import glosario
                data = glosario.guia(ACCIONES, META, {n: d for n, d in DEUDA.items() if not d.get("oculto")})
            elif ruta == "/api/dia":
                f = qs.get("fecha", [None])[0]
                if not f:
                    raise ValueError("Falta la fecha (AAAA-MM-DD)")
                reg = cached(("dia", f, date.today()), 600, lambda: dia(f))
                fmt = qs.get("formato", ["json"])[0]
                nombre = f"Alzea_analisis_{f}"
                if fmt == "pdf":
                    return self.send(reportes.pdf_dia(reg), "application/pdf", extra={"Content-Disposition": f'attachment; filename="{nombre}.pdf"'})
                if fmt == "xlsx":
                    return self.send(reportes.xlsx_dia(reg), XLSX, extra={"Content-Disposition": f'attachment; filename="{nombre}.xlsx"'})
                if fmt == "md":
                    return self.send(reportes.md_dia(reg).encode(), "text/markdown; charset=utf-8", extra={"Content-Disposition": f'attachment; filename="{nombre}.md"'})
                data = reg
            else:
                desde, hasta = periodo(qs)
                clave = (desde, hasta, date.today())
                if ruta == "/api/analisis":
                    data = analisis_completo(tickers, desde, hasta)
                elif ruta == "/api/deuda":
                    data = analisis_completo(list(ACCIONES)[:1], desde, hasta)["deuda"]
                elif ruta == "/api/comparar":
                    data = comparacion(tickers, desde, hasta, qs.get("moneda", ["mxn"])[0])
                elif ruta == "/api/simulacion":
                    data = cached(("sim", clave, qs.get("moneda", ["mxn"])[0]), 300, lambda: simulacion(desde, hasta, qs.get("moneda", ["mxn"])[0]))
                elif ruta == "/api/archivo":
                    data = cached(("arch", clave), 300, lambda: archivo(desde, hasta))
                elif ruta == "/api/contexto":
                    import contexto
                    data = contexto.construir(desde, hasta)
                elif ruta == "/api/informe":
                    fmt = qs.get("formato", ["json"])[0]
                    inf = cached(("inf", clave), 600, lambda: informe(desde, hasta))
                    nombre = f"Alzea_informe_{desde:%Y%m%d}_{hasta:%Y%m%d}"
                    if fmt == "pdf":
                        return self.send(reportes.pdf_periodo(inf), "application/pdf", extra={"Content-Disposition": f'attachment; filename="{nombre}.pdf"'})
                    if fmt == "xlsx":
                        return self.send(reportes.xlsx_periodo(inf), XLSX, extra={"Content-Disposition": f'attachment; filename="{nombre}.xlsx"'})
                    data = inf
                else:
                    return self.send(b'{"error":"ruta no encontrada"}', status=404)
            self.send(json.dumps(data, ensure_ascii=False).encode())
        except Exception as e:
            self.send(json.dumps({"error": str(e)}, ensure_ascii=False).encode(), status=502)


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


if __name__ == "__main__":
    print(f"Alzea en http://localhost:{PORT}  (Ctrl+C para detener)")
    ThreadingHTTPServer(("127.0.0.1", PORT), Handler).serve_forever()
