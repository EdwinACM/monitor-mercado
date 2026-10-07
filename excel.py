"""Libros de Excel en formato IPN: una hoja por emisora e instrumento, columnas con nombres claros y cifras con dos decimales.

Reglas de redacción: los textos van como frase (primera letra en mayúscula, el resto en minúscula salvo nombres propios y siglas),
cada hoja explica qué muestra y cada medida trae su significado, sin repetir información entre hojas.
"""
import io
import math
import re
from datetime import date, datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

import glosario
import reportes as R

TZ = ZoneInfo("America/Mexico_City")
G, GOSC, GSUAVE, GRIS, GRIS_S = "750946", "4F0630", "F7ECF1", "636569", "E8E9EA"
VERDE, ROJO = "CFEBDD", "F6D3CF"
ANCHOS = [42, 17, 17, 17, 17, 17, 17, 30]
NCOL = len(ANCHOS)
ANCHO_TOTAL = sum(ANCHOS)

N2 = "#,##0.00;-#,##0.00;0.00"
PCT = '+0.00"%";-0.00"%";0.00"%"'
PCT_S = '0.00"%"'
PB = '+0.00" pb";-0.00" pb";0.00" pb"'
MX = '"$"#,##0.00;-"$"#,##0.00'
US = '"US$"#,##0.00;-"US$"#,##0.00'
ENT = "+0;-0;0"
FECHA = "dd/mm/yyyy"
_fino = Side(style="thin", color=GRIS_S)


def frase(s):
    """Texto en forma de frase: si viene todo en mayúsculas lo pasa a minúsculas con la primera en mayúscula."""
    if not s:
        return s
    return s.capitalize() if s.isupper() and len(s) > 3 else s[0].upper() + s[1:]


def moneda_fmt(m):
    return US if m == "USD" else MX


def nombre_hoja(txt, usados):
    base = re.sub(r"[\[\]:*?/\\]", " ", txt).strip()[:31] or "Hoja"
    n, k = base, 2
    while n in usados:
        n = f"{base[:28]} {k}"
        k += 1
    usados.add(n)
    return n


def d_iso(s):
    return date.fromisoformat(s) if s else None


class Hoja:
    """Hoja con título institucional y bloques (texto, medidas con su significado, tablas)."""

    def __init__(self, wb, nombre, titulo, subtitulo, primera=False):
        self.ws = wb.active if primera else wb.create_sheet(nombre)
        if primera:
            self.ws.title = nombre
        self.nombre = self.ws.title
        ws = self.ws
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.orientation, ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = "landscape", 1, 0
        for j, w in enumerate(ANCHOS, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        for r in (1, 2):
            for c in range(1, NCOL + 1):
                ws.cell(r, c).fill = PatternFill("solid", fgColor=G)
        ws.row_dimensions[1].height = 26
        ws["A1"], ws["A2"] = titulo, subtitulo
        ws["A1"].font, ws["A2"].font = Font(bold=True, size=16, color="FFFFFF"), Font(size=10.5, color="FFFFFF")
        ws["A1"].alignment = ws["A2"].alignment = Alignment(vertical="center")
        self.r = 4

    # ---- bloques de texto
    def bloque(self, txt):
        self.r += 1
        ws = self.ws
        c = ws.cell(self.r, 1, txt)
        c.font = Font(bold=True, size=12, color=G)
        for j in range(1, NCOL + 1):
            ws.cell(self.r, j).border = Border(bottom=Side(style="medium", color=G))
        self.r += 1

    def parrafo(self, txt, italica=False, color="000000", alto_linea=15):
        ws = self.ws
        ws.merge_cells(start_row=self.r, start_column=1, end_row=self.r, end_column=NCOL)
        c = ws.cell(self.r, 1, txt)
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
        c.font = Font(italic=italica, color=color, size=10)
        ws.row_dimensions[self.r].height = alto_linea * max(1, math.ceil(len(txt) / (ANCHO_TOTAL * 1.05)))
        self.r += 1

    def nota(self, txt):
        self.parrafo(txt, italica=True, color=GRIS)

    def enlace(self, texto, url):
        c = self.ws.cell(self.r, 1, texto)
        c.hyperlink = url
        c.font = Font(color=G, underline="single", size=10)
        self.r += 1

    # ---- medidas con significado
    def medidas(self, items):
        """items: [(medida, valor, formato|None, significado)] en tres columnas: Medida, Valor y Qué significa."""
        ws = self.ws
        for j, t in enumerate(("Medida", "Valor", "Qué significa"), 1):
            c = ws.cell(self.r, j, t)
            c.fill, c.font = PatternFill("solid", fgColor=G), Font(bold=True, color="FFFFFF")
            c.alignment = Alignment(horizontal="center" if j > 1 else "left", vertical="center")
        for j in range(4, NCOL + 1):
            ws.cell(self.r, j).fill = PatternFill("solid", fgColor=G)
        ws.merge_cells(start_row=self.r, start_column=3, end_row=self.r, end_column=NCOL)
        self.r += 1
        for i, (med, val, fmt, sig) in enumerate(items):
            a, b = ws.cell(self.r, 1, med), ws.cell(self.r, 2, val)
            c = ws.cell(self.r, 3, sig)
            ws.merge_cells(start_row=self.r, start_column=3, end_row=self.r, end_column=NCOL)
            a.font = Font(bold=True, size=10)
            b.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True, indent=1)
            if fmt and isinstance(val, (int, float)):
                b.number_format = fmt
            a.alignment = Alignment(vertical="top", wrap_text=True)
            c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
            c.font = Font(size=9.5, color=GRIS)
            ws.row_dimensions[self.r].height = max(15, 13.5 * math.ceil(len(med) / (ANCHOS[0] * 0.95)), 13.5 * math.ceil(len(sig or "") / 110), 13.5 * math.ceil(len(str(val)) / 16) if isinstance(val, str) else 15)
            for j in range(1, NCOL + 1):
                ws.cell(self.r, j).border = Border(bottom=_fino)
                if i % 2:
                    ws.cell(self.r, j).fill = PatternFill("solid", fgColor=GSUAVE)
            self.r += 1

    # ---- tabla
    def tabla(self, cols, filas, formatos=None, colorear=None, fusion=()):
        """cols: encabezados; filas: listas (una celda puede ser (valor, formato)); fusion: [(col_ini, col_fin)] celdas que se combinan.
        Devuelve (fila del encabezado, última fila)."""
        ws = self.ws
        fus = {c1: c2 for c1, c2 in fusion}
        ocultas = {j for c1, c2 in fusion for j in range(c1 + 1, c2 + 1)}
        ancho_de = lambda j: sum(ANCHOS[j - 1:fus[j]]) if j in fus else ANCHOS[j - 1]
        enc = self.r
        for j in range(1, NCOL + 1):
            c = ws.cell(enc, j)
            c.fill, c.font = PatternFill("solid", fgColor=G), Font(bold=True, color="FFFFFF")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for j, t in zip(sorted(set(range(1, NCOL + 1)) - ocultas), cols):
            ws.cell(enc, j, t)
        for c1, c2 in fusion:
            ws.merge_cells(start_row=enc, start_column=c1, end_row=enc, end_column=c2)
        ws.row_dimensions[enc].height = max(24, 14.5 * max(math.ceil(len(str(t)) / max(6, ancho_de(j) * 0.95)) for j, t in zip(sorted(set(range(1, NCOL + 1)) - ocultas), cols)) + 6)
        self.r += 1
        for i, f in enumerate(filas):
            alto = 15
            posiciones = sorted(set(range(1, NCOL + 1)) - ocultas)
            for j in range(1, NCOL + 1):
                c = ws.cell(self.r, j)
                c.border = Border(bottom=_fino)
                if i % 2:
                    c.fill = PatternFill("solid", fgColor=GSUAVE)
            for j, v in zip(posiciones, f):
                fmt = (formatos or {}).get(j - 1)
                if isinstance(v, tuple):
                    v, fmt = v
                c = ws.cell(self.r, j, v)
                if fmt and isinstance(v, (int, float, date)):
                    c.number_format = fmt
                if isinstance(v, (int, float, date)) and not isinstance(v, bool):
                    c.alignment = Alignment(horizontal="right", vertical="top")
                else:
                    c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left", indent=1 if j > 1 else 0)
                    alto = max(alto, 13.5 * math.ceil(len(str(v)) / max(8, ancho_de(j) * 1.0)) + 1.5)
            for c1, c2 in fusion:
                ws.merge_cells(start_row=self.r, start_column=c1, end_row=self.r, end_column=c2)
            if colorear:
                for j, color in (colorear(f) or {}).items():
                    ws.cell(self.r, posiciones[j]).fill = PatternFill("solid", fgColor=color)
            ws.row_dimensions[self.r].height = alto
            self.r += 1
        return enc, self.r - 1


def _color_decision(f, col):
    """Verde para comprar, rojo para vender, sin color para mantener o esperar."""
    v = str(f[col]).lower()
    return {col: VERDE} if v.startswith("compra") else {col: ROJO} if v.startswith("vend") or v.startswith("venta") else {}


def _guardar(wb):
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _escudo(ws):
    try:
        img = XLImage(str(R.MARCA / "ipn-escudo-blanco.png"))
        img.width, img.height = 38, 56
        ws.add_image(img, f"{get_column_letter(NCOL)}1")
    except Exception:
        pass


SIGNIFICADOS = [
    ("Decisión", "Recomendación educativa que resume nueve señales técnicas con pesos recalculados cada día: comprar, mantener o vender. No es asesoría financiera."),
    ("Puntaje", "Suma de las nueve señales, de −100 a +100. Desde +20 se sugiere comprar; desde −20, vender; entre ambos, mantener."),
    ("Rendimiento", "Cambio porcentual del precio entre dos fechas, en la moneda de la emisora. Se muestra con signo: positivo es ganancia."),
    ("Puntos base (pb)", "Centésimas de punto porcentual: 100 pb equivalen a 1 punto porcentual. Se usan para cambios en tasas."),
    ("Millones de pesos (mdp)", "Unidad de los montos de deuda. 1,000 mdp equivalen a mil millones de pesos."),
    ("Puntaje z", "Cuántas desviaciones estándar se aleja la última tasa del promedio de las últimas 12 subastas. Arriba de +1 es alta; debajo de −1, baja."),
    ("Correlación", "De −1 a +1: mide si dos precios se mueven juntos (cerca de +1), en sentido contrario (cerca de −1) o sin relación (cerca de 0)."),
    ("Beta", "Cuánto tiende a moverse una acción por cada 1 % que se mueve el S&P 500. Mayor que 1 es más volátil que el mercado."),
    ("Fuente oficial y referencia", "Oficial: la publica la propia institución (Banxico, Tesoro de EE. UU., Reserva Federal de Nueva York, BLS, Cboe). Referencia: datos de mercado de Yahoo Finance, no oficiales."),
    ("Redondeo", "Todas las cifras se muestran con dos decimales; las celdas conservan la precisión original para que puedas calcular con ellas."),
]


def _leeme(wb, titulo, sub, parrafos, indice):
    h = Hoja(wb, "Léeme", titulo, sub, primera=True)
    _escudo(h.ws)
    h.bloque("Qué es este libro")
    for p in parrafos:
        h.parrafo(p)
    h.bloque("Qué contiene cada hoja")
    enc, fin = h.tabla(["Hoja", "Qué muestra"], [[n, d] for n, d in indice], fusion=[(2, NCOL)])
    for i, (n, _) in enumerate(indice):
        c = h.ws.cell(enc + 1 + i, 1)
        c.hyperlink = f"#'{n}'!A1"
        c.font = Font(color=G, underline="single", bold=True)
    h.bloque("Cómo leer las cifras")
    h.tabla(["Término", "Significado"], [[t, d] for t, d in SIGNIFICADOS], fusion=[(2, NCOL)])
    h.bloque("Aviso")
    h.parrafo(R.AVISO)
    h.nota(f"Elaborado el {R.fecha_corta(datetime.now(TZ).date().isoformat())} a las {datetime.now(TZ):%H:%M} (hora del centro de México).")
    return h


# ======================================================================= período
def libro_periodo(inf):
    an, ctx, arc, sim = inf["analisis"], inf["contexto"], inf["archivo"], inf["simulacion"]
    p = inf["periodo"]
    srv = __import__("server")
    acc = [x for x in an["acciones"] if "error" not in x]
    hist = R.historial(arc)
    gl = {a["ticker"]: a for a in glosario.guia({x["ticker"]: x["nombre"] for x in acc}, srv.META, {})["acciones"]}
    gd = {d["nombre"]: d for d in glosario.guia({}, {}, {n: {"codigo": "", "emisor": ""} for n in glosario.DEUDA})["deuda"]}
    usados = {"Léeme", "Resumen", "Entorno global", "Relación con EE. UU.", "Aciertos de las decisiones", "Simulación", "Fuentes"}
    hojas_acc = {x["ticker"]: nombre_hoja(x["ticker"].replace(".MX", ""), usados) for x in acc}
    hojas_deu = {x["nombre"]: nombre_hoja(frase(x["nombre"]), usados) for x in an["deuda"]}
    sub = f"Periodo del {R.fecha_corta(p['desde'])} al {R.fecha_corta(p['hasta'])} · último cierre {R.fecha_corta(p['ultimo_cierre'] or p['hasta'])}"

    wb = Workbook()
    indice = [("Resumen", "Lectura del periodo y tablero de cierres y decisiones de todas las emisoras y de la deuda."),
              ("Entorno global", "Dólar, tasas, inflación y bolsas del mundo, con su fuente y qué significa cada indicador."),
              ("Relación con EE. UU.", "Qué tan ligada está cada acción al S&P 500 y al dólar, y el efecto del tipo de cambio en las acciones de EE. UU.")]
    indice += [(hojas_acc[x["ticker"]], f"{x['nombre']}: cierre, decisión y por qué, estadística de riesgo, qué la mueve y cierres diarios con su gráfica.") for x in acc]
    indice += [(hojas_deu[x["nombre"]], f"{frase(x['nombre'])}: qué es, última cifra, decisión y resultados semana a semana con su gráfica.") for x in an["deuda"]]
    indice += [("Aciertos de las decisiones", "Cuántas veces acertó cada tipo de decisión y los últimos cambios de decisión."),
               ("Simulación", "Qué habría pasado al seguir las decisiones frente a comprar y mantener."),
               ("Fuentes", "De dónde vienen los datos y si la fuente es oficial.")]
    _leeme(wb, "Informe de mercados", sub,
           [f"Reúne el análisis de {len(acc)} acciones de la Bolsa Mexicana de Valores, el Nasdaq y la NYSE, y de {len(an['deuda'])} instrumentos de deuda, con el entorno global que puede mover sus precios. "
            "Cada emisora e instrumento tiene su propia hoja.",
            "El periodo siempre inicia el 1 de marzo de 2026 o después. Las decisiones se calculan con los datos disponibles hasta cada fecha, sin ver el futuro."], indice)

    # --- Resumen
    h = Hoja(wb, "Resumen", "Resumen del periodo", sub)
    h.bloque("Lectura del periodo")
    h.parrafo(an["mercado"].get("texto_periodo") or an["mercado"]["texto"])
    for t in ctx["lectura"][:3]:
        h.parrafo(t)
    h.bloque("Acciones: cierres y decisiones")
    filas = []
    for x in sorted(acc, key=lambda z: (z["pais"] != "México", z["nombre"])):
        m = moneda_fmt(x["moneda"])
        filas.append([x["nombre"], f"{x['mercado']} · {x['moneda']}", (x["serie"]["cierre"][0], m), (x["cierre"], m), (x["var_pct"], PCT), (x["ret_periodo"], PCT), x["decision"], (x["score"], ENT)])
    h.tabla(["Emisora", "Mercado y moneda", "Cierre inicial", "Cierre final", "Cambio del último día", "Rendimiento del periodo", "Decisión", "Puntaje"], filas, colorear=lambda f: _color_decision(f, 6))
    h.nota("Puntaje de −100 a +100. Los precios están en la moneda de cada mercado: pesos mexicanos (MXN) o dólares estadounidenses (USD).")
    h.bloque("Deuda: última cifra y decisión")
    filas = []
    for x in an["deuda"]:
        if x["valor"] is None:
            val = "Sin dato"
        elif x["tipo"] == "monto":
            val = (x["valor"] / 1000, N2)
        else:
            val = (x["valor"], N2)
        unidad = "Millones de pesos colocados" if x["tipo"] == "monto" else "Pesos por título" if x["tipo"] == "precio" else frase(x["unidad"])
        filas.append([frase(x["nombre"]), frase(x["emisor"]), unidad, val, (x.get("var_pb"), PB) if x.get("var_pb") is not None else "—", (x["z"], '+0.00;-0.00;0.00') if x.get("z") is not None else "—", frase(x["tendencia"]), x["decision"]])
    h.tabla(["Instrumento", "Emisor", "Unidad", "Última cifra", "Cambio frente a la anterior", "Puntaje z", "Tendencia", "Decisión"], filas, colorear=lambda f: _color_decision(f, 7))
    _enlazar_primera_columna(h.ws, {frase(x["nombre"]): hojas_deu[x["nombre"]] for x in an["deuda"]})
    _enlazar_primera_columna(h.ws, {x["nombre"]: hojas_acc[x["ticker"]] for x in acc})

    # --- Entorno global
    h = Hoja(wb, "Entorno global", "Entorno global", sub)
    h.nota("Indicadores que ayudan a explicar por qué suben o bajan las acciones y la deuda. Las tasas se miden en % anual; el cambio de una tasa se muestra en puntos base (pb).")
    filas = []
    ind = ctx["indicadores"]
    for x in ind:
        fm = PB if x["cambio_unidad"] == "pb" else PCT
        filas.append([x["nombre"], frase(x["grupo"]), (x["ini"]["valor"], N2), (x["fin"]["valor"], N2), (x["cambio"], fm), "Oficial" if x["fuente"]["oficial"] else "Referencia, no oficial",
                      x["fuente"]["nombre"].split(" · ")[0], frase(x["unidad"])])
    h.tabla(["Indicador", "Grupo", "Valor al inicio del periodo", "Valor al final", "Cambio en el periodo", "Tipo de fuente", "Fuente", "Unidad"], filas)
    h.bloque("Qué es cada indicador y cómo se relaciona con tus instrumentos")
    h.tabla(["Indicador", "Qué es", "Cómo se relaciona con tus instrumentos"], [[x["nombre"], x["que_es"], x["como_afecta"]] for x in ind], fusion=[(2, 4), (5, NCOL)])
    if ctx["diferenciales"]:
        h.bloque("Diferenciales de tasas")
        h.medidas([(d["nombre"], (d["pb"], PB)[0], PB, f"{d['detalle']}. {d['lectura']}") for d in ctx["diferenciales"]])

    # --- Relación con EE. UU.
    h = Hoja(wb, "Relación con EE. UU.", "Relación con EE. UU. y el dólar", sub)
    h.nota("Se calcula con los rendimientos diarios del periodo. Ninguna cifra prueba que un mercado cause el movimiento del otro; solo mide si se mueven juntos.")
    h.bloque("Sensibilidad de cada acción")
    h.tabla(["Emisora", "País", "Moneda", "Sesiones usadas", "Correlación con el S&P 500", "Beta frente al S&P 500", "Correlación con el dólar", "Correlación con el índice IPC"],
            [[s["nombre"], s["pais"], s["moneda"], (s["n"], "0"), (s["corr_sp500"], N2), (s["beta_sp500"], N2), (s["corr_dolar"], N2), (s["corr_ipc"], N2) if s["corr_ipc"] is not None else "No aplica"] for s in ctx["sensibilidades"]])
    h.nota("Correlación con el dólar: se usa el tipo de cambio de cierre de jornada de Banxico, que coincide con la hora de cierre de las bolsas. Si es negativa, la acción tiende a subir cuando el peso se aprecia.")
    if ctx["efecto_cambiario"]:
        h.bloque("Acciones de EE. UU. vistas desde México")
        h.tabla(["Acción", "Desde", "Hasta", "Rendimiento en dólares", "Variación del dólar frente al peso", "Rendimiento en pesos"],
                [[e["nombre"], (d_iso(e["desde"]), FECHA), (d_iso(e["hasta"]), FECHA), (e["ret_usd"], PCT), (e["ret_dolar"], PCT), (e["ret_pesos"], PCT)] for e in ctx["efecto_cambiario"]])
        h.nota("Un peso más débil (dólar más caro) suma al rendimiento en pesos; uno más fuerte lo resta. La variación del dólar usa el tipo de cambio de cierre de jornada, por eso difiere ligeramente de la del FIX que aparece en el entorno global.")

    # --- Una hoja por acción
    for x in acc:
        _hoja_accion(wb, hojas_acc[x["ticker"]], x, ctx, R.estadisticas(x), gl.get(x["ticker"], {}), hist.get(x["ticker"], []), sub)
    # --- Una hoja por instrumento de deuda
    for x in an["deuda"]:
        _hoja_deuda(wb, hojas_deu[x["nombre"]], x, ctx, gd.get(x["nombre"], {}), sub)

    # --- Aciertos
    h = Hoja(wb, "Aciertos de las decisiones", "¿Acertaron las decisiones?", sub)
    h.nota("Una decisión de comprar acierta si el precio subió a 5 sesiones; la de vender, si bajó; la de mantener, si se movió 2 % o menos. Las decisiones de los últimos 5 días aún no se pueden evaluar.")
    h.tabla(["Decisión", "Casos evaluados", "Aciertos", "Porcentaje de aciertos", "Rendimiento medio a 5 sesiones"],
            [[a["decision"], (a["casos"], "#,##0"), (a["aciertos"], "#,##0"), (a["pct"], PCT_S) if a["pct"] is not None else "—", (a["ret5_medio"], PCT) if a["ret5_medio"] is not None else "—"] for a in arc["aciertos"]])
    h.nota("Un porcentaje cercano a 50 % indica que la decisión no superó al azar en este periodo.")
    if arc["cambios"]:
        h.bloque("Últimos cambios de decisión")
        h.tabla(["Fecha", "Emisora", "Decisión anterior", "Decisión nueva"], [[(d_iso(c["fecha"]), FECHA), c["ticker"], c["de"].capitalize() if c["de"] else "Sin dato previo", c["a"].capitalize()] for c in arc["cambios"]])

    # --- Simulación
    if sim.get("fechas"):
        m = sim["metricas"]
        h = Hoja(wb, "Simulación", "Simulación: seguir las decisiones o comprar y mantener", sub)
        h.parrafo(f"Con ${sim['capital']:,.2f} repartidos en partes iguales entre {len(sim['por_emisora'])} emisoras (acciones de EE. UU. medidas en pesos con el tipo de cambio de Banxico). "
                  f"Se compra cuando el puntaje llega a +20 o más y se sale a efectivo con −20 o menos; la decisión se ejecuta al cierre siguiente y cada operación paga {sim['costo_pct']:.2f} % de comisión.")
        h.bloque("Resultado del periodo")
        h.tabla(["Medida", "Siguiendo las decisiones", "Comprar y mantener"],
                [["Rendimiento del periodo", (m["ret_estrategia"], PCT), (m["ret_comprar_mantener"], PCT)], ["Valor final en pesos", (sim["estrategia"][-1], MX), (sim["comprar_mantener"][-1], MX)],
                 ["Caída máxima", (m["caida_estrategia"], PCT), (m["caida_comprar_mantener"], PCT)], ["Operaciones", (m["operaciones"], "0"), "Una por emisora"]])
        h.bloque("Resultado por emisora")
        h.tabla(["Emisora", "Siguiendo las decisiones", "Comprar y mantener", "Operaciones", "Tiempo invertido"],
                [[next((x["nombre"] for x in acc if x["ticker"] == e["ticker"]), e["ticker"]), (e["ret_estrategia"], PCT), (e["ret_comprar_mantener"], PCT), (e["operaciones"], "0"), (e["tiempo_en_mercado"], PCT_S)] for e in sim["por_emisora"]])
        h.nota("Es una prueba dentro de la misma muestra, sin impuestos ni deslizamiento de precios. El resultado pasado no garantiza resultados futuros.")

    # --- Fuentes
    h = Hoja(wb, "Fuentes", "Fuentes de los datos", sub)
    filas = [[f["nombre"], "Oficial" if f["oficial"] else "Referencia, no oficial", f["url"]] for f in ctx["fuentes"]] + [
        ["Yahoo Finance · precios de acciones", "Referencia, no oficial", "https://finance.yahoo.com/"],
        ["SEC EDGAR · reportes de empresas de EE. UU.", "Oficial", "https://www.sec.gov/edgar"],
        ["Bolsa Mexicana de Valores · emisoras", "Oficial", "https://www.bmv.com.mx/"]]
    enc, fin = h.tabla(["Fuente", "Tipo", "Enlace"], filas, fusion=[(3, NCOL)])
    for i in range(len(filas)):
        c = h.ws.cell(enc + 1 + i, 3)
        c.hyperlink = c.value
        c.font = Font(color=G, underline="single")
    return _guardar(wb)


def _enlazar_primera_columna(ws, mapa):
    for row in ws.iter_rows(min_col=1, max_col=1):
        c = row[0]
        if isinstance(c.value, str) and c.value in mapa and not c.hyperlink:
            c.hyperlink = f"#'{mapa[c.value]}'!A1"
            c.font = Font(color=G, underline="single", bold=True)


def _grafica(h, cat_col, val_col, fila_ini, fila_fin, ancla_fila, titulo, valores, fmt_y=N2):
    ch = LineChart()
    ch.title = titulo
    ch.height, ch.width = 7.4, 27
    ch.legend = None
    ch.y_axis.number_format = fmt_y
    ch.x_axis.number_format = "dd/mm/yy"
    ch.x_axis.tickLblSkip = max(1, (fila_fin - fila_ini) // 8)
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    lo, hi = min(valores), max(valores)
    pad = (hi - lo) * 0.08 or abs(hi) * 0.02 or 1
    ch.y_axis.scaling.min, ch.y_axis.scaling.max = round(lo - pad, 2), round(hi + pad, 2)
    datos = Reference(h.ws, min_col=val_col, min_row=fila_ini - 1, max_row=fila_fin)
    ch.add_data(datos, titles_from_data=True)
    ch.set_categories(Reference(h.ws, min_col=cat_col, min_row=fila_ini, max_row=fila_fin))
    s = ch.series[0]
    s.graphicalProperties.line.solidFill = G
    s.graphicalProperties.line.width = 19000
    s.marker.symbol = "none"
    s.smooth = False
    h.ws.add_chart(ch, f"A{ancla_fila}")


def _hoja_accion(wb, nombre, x, ctx, e, g, hist, sub):
    mon = moneda_fmt(x["moneda"])
    h = Hoja(wb, nombre, f"{x['nombre']} ({x['ticker'].replace('.MX', '')})", f"{x['mercado']} · {'dólares estadounidenses' if x['moneda'] == 'USD' else 'pesos mexicanos'} · cierre del {R.fecha_corta(x['fecha'])}")
    s = x["stats"]
    h.bloque("Decisión del día")
    h.medidas([
        ("Decisión", x["decision"], None, "Recomendación educativa según nueve señales técnicas con pesos recalculados cada día. No es asesoría financiera."),
        ("Puntaje", (x["score"]), ENT, "De −100 (venta fuerte) a +100 (compra fuerte). Desde +20 se sugiere comprar; desde −20, vender; entre ambos, mantener."),
        ("Confianza", x["confianza"], None, "Qué tan coherentes son entre sí los criterios que respaldan la decisión."),
        ("Cierre", x["cierre"], mon, f"Precio de cierre de la última sesión ({R.fecha_corta(x['fecha'])})."),
        ("Cambio del último día", x["var_pct"], PCT, "Variación del precio frente al cierre anterior."),
        ("Rendimiento del periodo", x["ret_periodo"], PCT, "Cambio del precio entre el primer y el último cierre del periodo consultado."),
        ("Rendimiento a 5 sesiones", x["ret5"], PCT, "Cambio del precio en la última semana de operación."),
        ("Rendimiento a 20 sesiones", x["ret20"], PCT, "Cambio del precio en el último mes de operación."),
        ("Tendencia de 30 sesiones", x["tendencia"]["etiqueta"].capitalize(), None, "Dirección de la recta que mejor ajusta los últimos 30 cierres; “lateral” significa sin tendencia clara."),
    ])
    h.bloque("Precio de cierre del periodo")
    ancla = h.r
    h.r += 16
    h.bloque("Estadística de riesgo")
    medidas = [("Volatilidad anual", s["vol_anual"], PCT_S, "Qué tanto oscila el precio en un año; más alta significa más riesgo."),
               ("Caída máxima del periodo", s["max_drawdown"], PCT, "Mayor baja desde un máximo hasta el mínimo siguiente.")]
    if e:
        medidas += [("Pérdida diaria probable (VaR 95 %)", e["var95"], PCT_S, "En 95 de cada 100 sesiones del periodo la pérdida diaria no superó esta cifra."),
                    ("Pérdida promedio en las peores sesiones (CVaR 95 %)", e["cvar95"], PCT_S, "Promedio de las pérdidas diarias cuando se supera el VaR; mide qué tan malo puede ser lo peor."),
                    ("Rendimiento entre riesgo", s["ratio_rend_riesgo"], N2, "Rendimiento del periodo dividido entre su volatilidad; más alto es mejor."),
                    ("Días al alza", e["dias_alza"], PCT_S, "Porcentaje de sesiones en que el precio subió."),
                    ("Mejor día", e["mejor"][0], PCT, f"Mayor alza en una sola sesión; ocurrió el {R.fecha_corta(e['mejor'][1])}."),
                    ("Peor día", e["peor"][0], PCT, f"Mayor baja en una sola sesión; ocurrió el {R.fecha_corta(e['peor'][1])}."),
                    ("Máximo del periodo", e["maximo"][0], mon, f"Cierre más alto; ocurrió el {R.fecha_corta(e['maximo'][1])}."),
                    ("Mínimo del periodo", e["minimo"][0], mon, f"Cierre más bajo; ocurrió el {R.fecha_corta(e['minimo'][1])}.")]
    medidas += [("Fuerza relativa (RSI de 14 sesiones)", x["rsi"], N2, "De 0 a 100. Arriba de 70 indica posible sobrecompra; debajo de 30, posible sobreventa."),
                ("Soporte de 20 sesiones", s["soporte"], mon, "Precio mínimo reciente donde la baja se ha detenido."),
                ("Resistencia de 20 sesiones", s["resistencia"], mon, "Precio máximo reciente donde el alza se ha detenido.")]
    h.medidas(medidas)
    sens = next((z for z in ctx["sensibilidades"] if z["ticker"] == x["ticker"]), None)
    efe = next((z for z in ctx["efecto_cambiario"] if z["ticker"] == x["ticker"]), None)
    if sens:
        h.bloque("Relación con EE. UU. y el dólar")
        it = [("Correlación con el S&P 500", sens["corr_sp500"], N2, "De −1 a +1: si se mueve junto con la bolsa de EE. UU."),
              ("Beta frente al S&P 500", sens["beta_sp500"], N2, "Cuánto tiende a moverse por cada 1 % del S&P 500."),
              ("Correlación con el dólar", sens["corr_dolar"], N2, "Negativa: tiende a subir cuando el peso se aprecia.")]
        if sens["corr_ipc"] is not None:
            it.append(("Correlación con el IPC", sens["corr_ipc"], N2, "Si se mueve junto con la bolsa mexicana."))
        if efe:
            it += [("Rendimiento en dólares", efe["ret_usd"], PCT, "Lo que ganó o perdió la acción en su moneda."),
                   ("Variación del dólar frente al peso", efe["ret_dolar"], PCT, "Cuánto subió o bajó el dólar en el mismo lapso."),
                   ("Rendimiento en pesos", efe["ret_pesos"], PCT, "Lo que ganó o perdió quien invirtió desde México.")]
        h.medidas(it)
    h.bloque("Qué es y qué la mueve")
    h.parrafo(g.get("que_es", ""))
    h.parrafo("Qué la mueve. " + g.get("que_la_mueve", ""))
    fac = R._factores(x, ctx)
    if fac:
        h.bloque("Qué podría mover su precio ahora")
        for t in fac:
            h.parrafo("• " + t)
    h.bloque(f"Por qué la decisión es {x['decision'].lower()}")
    h.tabla(["Señal", "Puntos", "Correlación reciente", "Qué se observó"], [[c["criterio"], (c["puntos"], ENT), (c.get("correlacion"), '+0.00;-0.00;0.00') if c.get("correlacion") is not None else "—", c["detalle"]] for c in x["componentes"]], fusion=[(4, NCOL)])
    h.nota("Correlación reciente: qué tanto anticipó esa señal el rendimiento a 5 sesiones en las últimas 250 sesiones de las 10 emisoras. Con valor absoluto menor a 0.03 la señal no suma; si es negativa, se interpreta al revés.")
    h.parrafo(x["accion"])
    h.bloque("Respaldo estadístico y riesgos")
    h.parrafo(x["respaldo"])
    for t_ in x["riesgos"]:
        h.parrafo("• " + t_)
    h.bloque("Sugerencia")
    h.parrafo(R._sugerencia(x, e))
    b = x["backtest"]
    if b["n_compra"] >= 8 or b["n_venta"] >= 8:
        h.nota(f"Prueba histórica a {b['horizonte']} sesiones con {b['muestra']} sesiones: comprar acertó {b['aciertos_compra']:.2f} % en {b['n_compra']} casos y vender {b['aciertos_venta']:.2f} % en {b['n_venta']}. "
               f"Rendimiento medio de cualquier sesión: {b['base_rend_medio']:+.2f} %. Es una prueba dentro de la misma muestra: calibra la confianza, no la garantiza.")
    if (g.get("fuente") or {}).get("nombre"):
        h.enlace(f"Fuente oficial: {g['fuente']['nombre']}", g["fuente"]["url"])
    h.nota("Precios: Yahoo Finance (referencia, no oficial).")
    h.bloque("Cierres diarios y decisión de cada sesión")
    h.nota("De la sesión más antigua a la más reciente. Rendimiento a 5 sesiones: lo que cambió el precio cinco sesiones después; ¿Acertó? compara la decisión con ese resultado.")
    filas = [[(d_iso(r["fecha"]), FECHA), (r["cierre"], mon), (r["var_pct"], PCT) if r["var_pct"] is not None else "—", r["decision"], (r["s"], ENT),
              (r["r5"], PCT) if r["r5"] is not None else "Por evaluar", {True: "Sí", False: "No", None: "—"}[r["a"]]] for r in hist]
    if filas:
        ini, fin = h.tabla(["Fecha", "Cierre", "Cambio del día", "Decisión", "Puntaje", "Rendimiento a 5 sesiones", "¿Acertó?"], filas, colorear=lambda f: _color_decision(f, 3))
        _grafica(h, 1, 2, ini + 1, fin, ancla, f"Cierre diario ({x['moneda']})", [r["cierre"] for r in hist])
    h.ws.freeze_panes = "A3"


def _hoja_deuda(wb, nombre, x, ctx, g, sub):
    tipo = x["tipo"]
    nom = frase(x["nombre"])
    h = Hoja(wb, nombre, nom, f"{frase(x['emisor'])} · {frase(x['unidad'])}")
    fmt_v = N2
    if x["valor"] is None:
        ult = "Sin dato"
    elif tipo == "monto":
        ult = x["valor"] / 1000
    else:
        ult = x["valor"]
    cifra = {"tasa": "Rendimiento o tasa de la última subasta, en % anual (o puntos porcentuales si así lo indica la unidad).", "precio": "Precio de la última subasta, en pesos por título.",
             "monto": "Millones de pesos colocados en el último mes o semana con datos."}[tipo]
    h.bloque("Decisión")
    it = [("Decisión", x["decision"], None, "Lectura educativa que compara la última cifra con las 12 anteriores. No es asesoría financiera."),
          (f"Última cifra ({R.fecha_corta(x['fecha'])})", ult, fmt_v, cifra),
          ("Tendencia", frase(x["tendencia"]), None, "Dirección de la recta que ajusta las últimas 8 cifras.")]
    if x.get("var_pb") is not None:
        it.append(("Cambio frente a la cifra anterior", x["var_pb"], PB, "Diferencia con la subasta o semana previa, en puntos base."))
    if x.get("z") is not None:
        it.append(("Puntaje z", x["z"], '+0.00;-0.00;0.00', "Distancia de la última cifra al promedio de las últimas 12, en desviaciones estándar."))
    h.medidas(it)
    h.bloque("Evolución")
    ancla = h.r
    h.r += 16
    h.bloque("Qué es y qué lo mueve")
    h.parrafo(g.get("que_es", ""))
    h.parrafo("Qué lo mueve. " + g.get("que_lo_mueve", ""))
    fac = R.factores_deuda(ctx)
    if fac:
        h.bloque("Qué podría mover su rendimiento ahora")
        for t in fac[:5]:
            h.parrafo("• " + t)
    h.bloque("Lectura")
    for q in x["por_que"]:
        h.parrafo("• " + q)
    if (g.get("fuente") or {}).get("nombre"):
        h.enlace(f"Fuente oficial: {g['fuente']['nombre']}", g["fuente"]["url"])
    if x["nombre"] in ("PAPEL COMERCIAL", "CERTIFICADOS BURSÁTILES") and ctx["privados"]:
        h.bloque("Detalle mensual del mercado privado (Banxico)")
        h.nota("Montos en millones de pesos. Una tasa de 0.00 significa que no hubo colocaciones de ese instrumento en el mes.")
        h.tabla(["Mes", "Tasa de certificados a corto plazo", "Tasa de certificados a mediano plazo", "Colocado a corto plazo", "Colocado en papel comercial", "Colocado a mediano y largo plazo"],
                [[(d_iso(m["mes"]), "mmm yyyy"), (m["tasa_cb_cp"], PCT_S), (m["tasa_cb_mp"], PCT_S), ((m["col_cp"] or 0) / 1000, N2), ((m["col_pc"] or 0) / 1000, N2), ((m["col_mlp"] or 0) / 1000, N2)] for m in ctx["privados"]])
    h.bloque("Resultados semana a semana")
    ex = x["extras"]
    cols = ["Fecha", {"tasa": "Tasa", "precio": "Precio", "monto": "Colocado en millones de pesos"}[tipo]] + (["Cambio frente a la anterior"] if tipo == "tasa" else []) + [e_ for e_ in ex]
    filas = []
    for d in x["datos"]:
        v = d["valor"] / 1000 if tipo == "monto" else d["valor"]
        fila = [(d_iso(d["fecha"]), FECHA), (v, N2)]
        if tipo == "tasa":
            fila.append((d["var_pb"], PB) if d.get("var_pb") is not None else "—")
        fila += [((d.get("extra") or {}).get(e_), N2) if (d.get("extra") or {}).get(e_) is not None else "—" for e_ in ex]
        filas.append(fila)
    if filas:
        ini, fin = h.tabla(cols, filas)
        if len(filas) >= 3:
            _grafica(h, 1, 2, ini + 1, fin, ancla, {"tasa": "Tasa por subasta (%)", "precio": "Precio por subasta", "monto": "Monto colocado (millones de pesos)"}[tipo], [(d["valor"] / 1000 if tipo == "monto" else d["valor"]) for d in x["datos"]])
    h.ws.freeze_panes = "A3"


# ============================================================================ día
def libro_dia(reg):
    usados = {"Léeme", "Resumen"}
    hojas = {e["ticker"]: nombre_hoja(e["ticker"], usados) for e in reg["emisoras"]}
    hojas_d = {x["nombre"]: nombre_hoja(frase(x["nombre"]), usados) for x in reg["deuda"]}
    sub = R.fecha_larga(reg["fecha"])
    wb = Workbook()
    indice = [("Resumen", "Lectura del día, entorno global y tablero de decisiones.")]
    indice += [(hojas[e["ticker"]], f"{e['nombre']}: cierre, decisión, por qué, qué hacer y qué pasó después.") for e in reg["emisoras"]]
    indice += [(hojas_d[x["nombre"]], f"{frase(x['nombre'])}: última cifra y decisión.") for x in reg["deuda"]]
    _leeme(wb, "Análisis del día", sub, ["Decisión de cada acción e instrumento de deuda al cierre de la sesión, calculada solo con datos disponibles hasta ese día, con lo que ocurrió después cuando ya hay datos."], indice)
    h = Hoja(wb, "Resumen", "Resumen del día", sub)
    h.bloque("Lectura del mercado")
    h.parrafo(reg["mercado"]["texto"])
    ent = (reg.get("entorno") or {}).get("indicadores") or []
    if ent:
        h.bloque("Entorno global del día")
        h.tabla(["Indicador", "Fecha del dato", "Valor", "Cambio frente a la sesión previa", "Tipo de fuente"],
                [[x["nombre"], (d_iso(x["fecha"]), FECHA), (x["valor"], N2), (x["cambio"], PB if x["cambio_unidad"] == "pb" else PCT), "Oficial" if x["oficial"] else "Referencia, no oficial"] for x in ent])
    h.bloque("Decisiones del día")
    filas = [[e["nombre"], e.get("mercado"), e.get("moneda"), (e["cierre"], moneda_fmt(e.get("moneda"))), (e["var_pct"], PCT), e["decision"], (e["score"], ENT), e["confianza"]] for e in reg["emisoras"]]
    h.tabla(["Emisora", "Mercado", "Moneda", "Cierre", "Cambio del día", "Decisión", "Puntaje", "Confianza"], filas, colorear=lambda f: _color_decision(f, 5))
    _enlazar_primera_columna(h.ws, {e["nombre"]: hojas[e["ticker"]] for e in reg["emisoras"]})
    h.bloque("Deuda")
    h.tabla(["Instrumento", "Emisor", "Fecha de la cifra", "Cifra", "Cambio (pb)", "Puntaje z", "Tendencia", "Decisión"],
            [[frase(x["nombre"]), frase(x["emisor"]), (d_iso(x["fecha"]), FECHA), (x["valor"], N2) if x["valor"] is not None else "—", (x["var_pb"], PB) if x.get("var_pb") is not None else "—",
              (x["z"], '+0.00;-0.00;0.00') if x.get("z") is not None else "—", frase(x["tendencia"]), x["decision"]] for x in reg["deuda"]], colorear=lambda f: _color_decision(f, 7))
    _enlazar_primera_columna(h.ws, {frase(x["nombre"]): hojas_d[x["nombre"]] for x in reg["deuda"]})
    for e in reg["emisoras"]:
        mon = moneda_fmt(e.get("moneda"))
        s = Hoja(wb, hojas[e["ticker"]], f"{e['nombre']} ({e['ticker']})", f"{e.get('mercado')} · {e.get('moneda')} · {sub}")
        s.bloque("Decisión")
        s.medidas([("Decisión", e["decision"], None, "Recomendación educativa según nueve señales técnicas con pesos recalculados cada día. No es asesoría financiera."),
                   ("Puntaje", e["score"], ENT, "De −100 a +100. Desde +20 comprar; desde −20 vender; entre ambos, mantener."),
                   ("Confianza", e["confianza"], None, "Qué tan coherentes son entre sí los criterios."),
                   ("Cierre", e["cierre"], mon, "Precio de cierre de la sesión."),
                   ("Cambio del día", e["var_pct"], PCT, "Variación frente al cierre anterior."),
                   ("Rendimiento a 5 sesiones previas", e["ret5"], PCT, "Cambio del precio en la última semana."),
                   ("Rendimiento a 20 sesiones previas", e["ret20"], PCT, "Cambio del precio en el último mes."),
                   ("Fuerza relativa (RSI de 14 sesiones)", e["rsi"], N2, "De 0 a 100. Arriba de 70, posible sobrecompra; debajo de 30, posible sobreventa."),
                   ("¿Cambió la decisión?", "Sí" if e.get("cambio") else "No", None, f"La sesión anterior era: {(e.get('veredicto_ant') or 'sin dato').capitalize()}.")])
        s.bloque("Por qué")
        s.tabla(["Criterio", "Puntos"], [[c[0], (c[1], ENT)] for c in e["componentes"]])
        for q in e["por_que"]:
            s.parrafo("• " + q)
        s.bloque("Qué hacer")
        s.parrafo(e["accion"])
        if e.get("respaldo"):
            s.bloque("Respaldo estadístico y riesgos")
            s.parrafo(e["respaldo"])
            for t_ in e.get("riesgos", []):
                s.parrafo("• " + t_)
        s.bloque("Qué pasó después")
        s.parrafo(R.posterior(e))
    for x in reg["deuda"]:
        s = Hoja(wb, hojas_d[x["nombre"]], frase(x["nombre"]), f"{frase(x['emisor'])} · {sub}")
        s.bloque("Decisión")
        s.medidas([("Decisión", x["decision"], None, "Lectura educativa que compara la última cifra con las 12 anteriores."),
                   (f"Cifra del {R.fecha_corta(x['fecha'])}", x["valor"] if x["valor"] is not None else "Sin dato", N2, "Última tasa, precio o monto publicado hasta esa fecha."),
                   ("Cambio frente a la anterior", x.get("var_pb") if x.get("var_pb") is not None else "—", PB, "Diferencia con la cifra previa, en puntos base."),
                   ("Puntaje z", x["z"] if x.get("z") is not None else "—", '+0.00;-0.00;0.00', "Distancia al promedio de las últimas 12, en desviaciones estándar.")])
        s.bloque("Por qué")
        for q in x["por_que"]:
            s.parrafo("• " + q)
    return _guardar(wb)
