"""Guarda el análisis diario en archivos (static/data/).

    python analisis_diario.py                      # agrega los días hábiles que falten (y refresca el último)
    python analisis_diario.py --desde 2026-03-02   # reconstruye desde esa fecha

Lo ejecuta GitHub Actions cada día hábil después del cierre (ver .github/workflows).
Archivos generados:
    analisis_diario.json    lo lee la sección «Archivo» de la app
    analisis_acciones.csv   una fila por fecha y emisora, con la decisión y sus razones
    analisis_deuda.csv      una fila por fecha e instrumento
    diario/AAAA-MM-DD.md    reporte legible de cada día: decisión, por qué y qué pasó después
Criterio de acierto (a 5 sesiones): comprar acierta si el precio subió; vender, si bajó; mantener, si se movió 2 % o menos.
"""
import argparse
import csv
import json
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import diario
import reportes
import server

OUT = server.STATIC / "data"
INICIO = server.INICIO.isoformat()

COLS_ACC = ["fecha", "ticker", "nombre", "mercado", "moneda", "cierre", "var", "var_pct", "ret5", "ret20", "rsi", "macd_hist", "pctb", "tendencia",
            "pendiente", "r2", "vol_anual", "max_drawdown", "score", "veredicto", "decision", "confianza", "veredicto_ant",
            "cambio", "razones", "accion", "ret_5", "ret_10", "acierto"]
COLS_DEU = ["fecha", "nombre", "emisor", "fecha_subasta", "valor", "var_pb", "z", "pendiente", "tendencia", "senal", "decision", "razones"]


def filas(registros):
    acc, deu = [], []
    for r in registros:
        for e in r["emisoras"]:
            f = {"fecha": r["fecha"], **{k: e.get(k) for k in COLS_ACC[1:]}}
            f["razones"] = " | ".join(e["por_que"])
            f["cambio"] = "Sí" if e.get("cambio") else "No"
            f["acierto"] = {True: "Sí", False: "No", None: ""}[e.get("acierto")]
            acc.append(f)
        for d in r["deuda"]:
            deu.append({"fecha": r["fecha"], "fecha_subasta": d["fecha"], "razones": " | ".join(d["por_que"]),
                        **{k: d.get(k) for k in COLS_DEU if k not in ("fecha", "fecha_subasta", "razones")}})
    return acc, deu


ROTULOS = {"fecha": "Fecha", "ticker": "Clave", "nombre": "Nombre", "mercado": "Mercado", "moneda": "Moneda", "cierre": "Cierre", "var": "Cambio del día",
           "var_pct": "Cambio del día (%)", "ret5": "Rendimiento previo a 5 sesiones (%)", "ret20": "Rendimiento previo a 20 sesiones (%)", "rsi": "Fuerza relativa (RSI)",
           "macd_hist": "Histograma MACD", "pctb": "Posición en las bandas de Bollinger", "tendencia": "Tendencia", "pendiente": "Pendiente de la tendencia",
           "r2": "Ajuste de la tendencia (R²)", "vol_anual": "Volatilidad anual (%)", "max_drawdown": "Caída máxima (%)", "score": "Puntaje", "veredicto": "Señal",
           "decision": "Decisión", "confianza": "Confianza", "veredicto_ant": "Señal anterior", "cambio": "¿Cambió la decisión?", "razones": "Por qué",
           "accion": "Qué hacer", "ret_5": "Rendimiento posterior a 5 sesiones (%)", "ret_10": "Rendimiento posterior a 10 sesiones (%)", "acierto": "¿Acertó?",
           "emisor": "Emisor", "fecha_subasta": "Fecha de la cifra", "valor": "Cifra", "var_pb": "Cambio (puntos base)", "z": "Puntaje z", "senal": "Señal"}


def escribir_csv(nombre, cols, filas_):
    """Tabla plana con encabezados legibles y números con dos decimales."""
    def v(x):
        return f"{x:.2f}" if isinstance(x, float) else x
    with open(OUT / nombre, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow([ROTULOS.get(c, c) for c in cols])
        for r in filas_:
            w.writerow([v(r.get(c)) for c in cols])


def escribir_xlsx(registros, acc, deu):
    G, hf = PatternFill("solid", fgColor="750946"), Font(bold=True, color="FFFFFF")  # guinda IPN
    verde, rojo = PatternFill("solid", fgColor="CFEBDD"), PatternFill("solid", fgColor="F6D3CF")
    wb = Workbook()

    def hoja(ws, cols, datos, anchos, wrap=()):
        ws.append(cols)
        for c in ws[1]:
            c.fill, c.font, c.alignment = G, hf, Alignment(horizontal="center", vertical="center", wrap_text=True)
        for f in datos:
            ws.append([f.get(c) for c in cols])
        for i, w in enumerate(anchos):
            ws.column_dimensions[get_column_letter(i + 1)].width = w
        for fila in ws.iter_rows(min_row=2):
            for j in wrap:
                fila[j].alignment = Alignment(wrap_text=True, vertical="top", horizontal="justify")
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Resumen diario"
    hoja(ws, ["fecha", "suben", "bajan", "puntaje_promedio", "sesgo", "texto"],
         [{"fecha": r["fecha"], **{k: r["mercado"].get(k) for k in ("suben", "bajan", "puntaje_promedio", "sesgo", "texto")}} for r in registros],
         [12, 8, 8, 16, 10, 120])
    ws = wb.create_sheet("Decisiones")
    anchos = {"fecha": 12, "ticker": 11, "nombre": 30, "mercado": 9, "moneda": 8, "razones": 100, "accion": 60, "decision": 22}
    hoja(ws, COLS_ACC, acc, [anchos.get(c, 10) for c in COLS_ACC], wrap=(COLS_ACC.index("razones"), COLS_ACC.index("accion")))
    for fila in ws.iter_rows(min_row=2):
        d = fila[COLS_ACC.index("decision")]
        d.fill = verde if "Comprar" in str(d.value) else rojo if "Vender" in str(d.value) else PatternFill()
        for c in ("cierre", "var", "var_pct", "ret5", "ret20", "rsi", "macd_hist", "pctb", "pendiente", "r2", "vol_anual", "max_drawdown", "ret_5", "ret_10"):
            fila[COLS_ACC.index(c)].number_format = "0.00"
    ws = wb.create_sheet("Deuda")
    hoja(ws, COLS_DEU, deu, [12, 20, 18, 14, 10, 10, 8, 11, 11, 11, 22, 100], wrap=(11,))
    wb.save(OUT / "analisis_diario.xlsx")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--desde", help="reconstruir desde esta fecha (AAAA-MM-DD)")
    args = ap.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)

    existente = {}
    ruta = OUT / "analisis_diario.json"
    if ruta.exists() and not args.desde:
        existente = {r["fecha"]: r for r in json.loads(ruta.read_text(encoding="utf-8"))["registros"]}

    accs = server.acciones(list(server.ACCIONES), server.ventana_historial())
    hist = {a["ticker"]: a["historico"] for a in accs if a["historico"]}
    if not hist:
        raise SystemExit("[ERROR] Yahoo Finance no devolvió datos; no se modificó ningún archivo.")
    deu = server.deuda(date.today() - timedelta(days=300))

    habiles = sorted({r["fecha"] for rows in hist.values() for r in rows})
    if args.desde:
        pendientes = [f for f in habiles if f >= args.desde]
    elif existente:
        ultimo = max(existente)
        pendientes = [f for f in habiles if f >= ultimo]
    else:
        pendientes = [f for f in habiles if f >= INICIO]

    nuevos = 0
    for f in pendientes:
        reg = diario.registro_dia(f, hist, deu, server.ACCIONES, server.META)
        if reg:
            existente[f] = reg
            nuevos += 1
    registros = diario.completar([existente[k] for k in sorted(existente)], hist)
    (OUT / "diario").mkdir(parents=True, exist_ok=True)
    for reg in registros:
        (OUT / "diario" / f"{reg['fecha']}.md").write_text(reportes.md_dia(reg), encoding="utf-8")
    ruta.write_text(json.dumps({"actualizado": datetime.now(server.TZ).isoformat(), "aviso": reportes.AVISO, "registros": registros},
                               ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    acc, deud = filas(registros)
    escribir_csv("analisis_acciones.csv", COLS_ACC, acc)
    escribir_csv("analisis_deuda.csv", COLS_DEU, deud)
    print(f"[OK] {nuevos} día(s) procesados; {len(registros)} en total ({registros[0]['fecha']} -> {registros[-1]['fecha']}).")


if __name__ == "__main__":
    main()
